from openpyxl import load_workbook

class ExcelProcessor:
    #Contructeur de la class ExcelProcessor.
    #Emmagasine le nom du dossier excel entrant et sortant 
    #ainsi que la classe SentimentAnalyzer.

    def __init__(self, input_file, output_file, analyzer):
        self.input_file = input_file
        self.output_file = output_file
        self.analyzer = analyzer

    def load_workbook(self):
        #Ouvre le dossier excel avec le chemin et retourne la première page et le dossier.

        workbook = load_workbook(self.input_file)
        sheet = workbook.active

        return workbook, sheet
    
    def process_comments(self, sheet):
        #lis les commentaires de la column B et utilise la classe SentimentAnalyzer pour obtenir le resultat de -1 a 1.
        # Mets ensuite les resultats dans la column D.

        for row in range(2, sheet.max_row + 1):
            comment = sheet.cell(row=row, column=2).value

            score = self.analyzer.calculate_score(str(comment))

            sheet.cell(row=row, column=4).value = round(score, 3)

    def save_workbook(self, workbook):
        #Sauvegarde les modifications dans un nouveau dossier xlsx.
        workbook.save(self.output_file)